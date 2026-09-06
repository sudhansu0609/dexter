"""
Folder ingestion: the files under the owner's knowledge roots become memory
document chunks Dexter can recall by meaning, with the file cited (roadmap v5,
4.1).

Roots come from DEXTER_KNOWLEDGE_ROOTS as `tag=path;tag=path`, e.g.

    notes=B:\\Notes;docs=C:\\Users\\me\\Documents;youtube=B:\\youtubeProjects

Each file is fingerprinted (mtime + size); unchanged files are skipped, changed
ones are re-chunked and re-embedded, vanished ones are removed. A nightly job
and a 60-second incremental pass at startup keep the index current; the
`reindex_knowledge` tool forces a full pass. Text families, PDF (pypdf) and
DOCX (python-docx) are supported; scanned PDFs with no text are recorded as
`empty` and skipped.
"""

import fnmatch
import hashlib
import os
import time
from typing import Any, Dict, Iterable, List, Optional, Tuple

import env_boot  # noqa: F401  - loads .env for every entry point

CHUNK_CHARS = 1200
MAX_FILE_BYTES = 10 * 1024 * 1024
MIN_TEXT_CHARS = 200
MIN_JSON_CHARS = 40   # a rendered studio record is short but still a real memory

TEXT_EXTENSIONS = (".md", ".markdown", ".txt", ".csv", ".rst")
PDF_EXTENSIONS = (".pdf",)
DOCX_EXTENSIONS = (".docx",)
JSON_EXTENSIONS = (".json",)  # only under the `studio` tag (structured records)

EXCLUDE_DIRS = {
    "node_modules", ".git", "__pycache__", "dist", "build", "target", ".venv", "venv", "voice_env",
    "static", "assets", ".cache", ".idea", ".vscode", "logs", "Downloads", "ComfyUI", "Adobe",
    "My Kindle Content", "Cline", "AirDroid", "AirDroidCast", "models", "buzzbrain", "data",
    # engine/tool caches and asset packs that carry text files but no knowledge
    "Saved", "Intermediate", "DerivedDataCache", "chrome-profile", "Wondershare", "Fonts",
    "Vectors", "oldThumbnails", "site-packages", "Lib", "Include", "Scripts",
}
# Persona prompts are noise everywhere except through the studio include list.
EXCLUDE_DIRS_NON_STUDIO = {"prompts"}
EXCLUDE_DIR_GLOBS = ("AirDroid*", ".*", "Python-*", "python-*", "webcache*", "*.egg-info", "*_cache")
EXCLUDE_FILE_GLOBS = (
    "package-lock.json", "*.min.*", "*.lock", "*.map", "LICENSE*", "license*", "COPYING*",
    "*.LICENSE.txt", "version.md", "DDCKey*", "CHANGELOG*", "requirements*.txt", "robots.txt",
)
# Inside a code repository (a folder with one of these markers) the docs and
# test trees describe the code, not the owner's life; README files still count.
CODE_REPO_MARKERS = ("package.json", "requirements.txt", "Cargo.toml", "pyproject.toml", "setup.py", "pyvenv.cfg")
CODE_REPO_SKIP_DIRS = {"docs", "doc", "tests", "test", "examples", "example", "spec", "specs"}

# What the `studio` root contributes (relative globs); everything else there is skipped.
STUDIO_INCLUDE = (
    "prompts/channels/*.md",
    "knowledge/saved_topics.json",
    "knowledge/agent_*_memory.json",
    "projects/*/*.md",
    "projects/*/*.json",
    "projects/*/*/*.md",
    "projects/*/*/*.json",
)

# Hot-patched by settings_store; read fresh on every scan.
ROOTS_SPEC = os.getenv("DEXTER_KNOWLEDGE_ROOTS", "")


def parse_roots(spec: Optional[str] = None) -> List[Tuple[str, str]]:
    """`tag=path;tag=path` -> [(tag, absolute path)], keeping only existing folders."""
    raw = spec if spec is not None else (os.getenv("DEXTER_KNOWLEDGE_ROOTS") or ROOTS_SPEC)
    roots: List[Tuple[str, str]] = []
    for part in (raw or "").split(";"):
        part = part.strip()
        if not part:
            continue
        if "=" in part:
            tag, path = part.split("=", 1)
        else:
            tag, path = os.path.basename(part.rstrip("\\/")) or "docs", part
        tag = tag.strip().lower() or "docs"
        path = os.path.realpath(os.path.expanduser(path.strip().strip('"')))
        if os.path.isdir(path):
            roots.append((tag, path))
    return roots


def root_paths() -> List[str]:
    return [path for _, path in parse_roots()]


# ───────────────────────── scanning ─────────────────────────


def _is_code_repo(base: str) -> bool:
    return any(os.path.exists(os.path.join(base, marker)) for marker in CODE_REPO_MARKERS)


def _excluded_dir(name: str, tag: str = "") -> bool:
    if name in EXCLUDE_DIRS:
        return True
    if tag != "studio" and name in EXCLUDE_DIRS_NON_STUDIO:
        return True
    return any(fnmatch.fnmatch(name, pattern) for pattern in EXCLUDE_DIR_GLOBS)


def _excluded_file(name: str) -> bool:
    return any(fnmatch.fnmatch(name, pattern) for pattern in EXCLUDE_FILE_GLOBS)


def _studio_wanted(root: str, path: str) -> bool:
    rel = os.path.relpath(path, root).replace("\\", "/")
    return any(fnmatch.fnmatch(rel, pattern) for pattern in STUDIO_INCLUDE)


def scan(roots: Optional[Iterable[Tuple[str, str]]] = None) -> List[Tuple[str, str]]:
    """Every ingestible file under the roots: [(path, tag)]."""
    found: List[Tuple[str, str]] = []
    roots = list(roots if roots is not None else parse_roots())
    all_roots = {os.path.normcase(path) for _, path in roots}
    for tag, root in roots:
        for base, dirs, files in os.walk(root):
            # A root nested inside another (the studio inside the youtube tree)
            # is walked once, under its own tag.
            in_repo = tag != "studio" and _is_code_repo(base)
            if os.path.exists(os.path.join(base, "pyvenv.cfg")):
                dirs[:] = []  # a virtual environment: nothing of the owner's in here
                continue
            dirs[:] = [
                d for d in dirs
                if not _excluded_dir(d, tag)
                and not (in_repo and d.lower() in CODE_REPO_SKIP_DIRS)
                and os.path.normcase(os.path.join(base, d)) not in all_roots
            ]
            for name in files:
                if _excluded_file(name):
                    continue
                lower = name.lower()
                path = os.path.join(base, name)
                if tag == "studio":
                    if not _studio_wanted(root, path):
                        continue
                elif lower.endswith(JSON_EXTENSIONS):
                    continue  # code-repo JSON is noise outside the studio root
                if not lower.endswith(TEXT_EXTENSIONS + PDF_EXTENSIONS + DOCX_EXTENSIONS + JSON_EXTENSIONS):
                    continue
                try:
                    if os.path.getsize(path) > MAX_FILE_BYTES:
                        continue
                except OSError:
                    continue
                found.append((path, tag))
    return found


# ───────────────────────── extraction ─────────────────────────


def _read_text(path: str) -> str:
    with open(path, "r", encoding="utf-8", errors="ignore") as handle:
        return handle.read()


def _read_pdf(path: str) -> str:
    from pypdf import PdfReader  # type: ignore

    reader = PdfReader(path)
    pages = []
    for number, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        text = text.strip()
        if text:
            pages.append(f"[p.{number}]\n{text}")
    return "\n\n".join(pages)


def _read_docx(path: str) -> str:
    import docx  # type: ignore  (python-docx)

    document = docx.Document(path)
    parts = [p.text for p in document.paragraphs if p.text and p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)


def _render_json(path: str, text: str) -> str:
    """Studio records as prose: one paragraph per topic / memory item."""
    import json

    try:
        data = json.loads(text)
    except Exception:
        return text
    records: List[Any]
    if isinstance(data, dict) and isinstance(data.get("topics"), list):
        records = data["topics"]
    elif isinstance(data, list):
        records = data
    elif isinstance(data, dict):
        records = [data]
    else:
        return text
    paragraphs = []
    for record in records:
        if not isinstance(record, dict):
            paragraphs.append(str(record))
            continue
        content = record.get("content")
        if isinstance(content, dict) and "user" in content:
            paragraphs.append(
                f"Exchange with {content.get('agent', 'an agent')} for {content.get('channel', '')}: "
                f"asked '{content.get('user', '')}'; reply: {content.get('reply', '')}"
            )
            continue
        bits = []
        for key, value in record.items():
            if value in (None, "", [], {}):
                continue
            if isinstance(value, (list, tuple)):
                value = ", ".join(str(v) for v in value)
            elif isinstance(value, dict):
                value = "; ".join(f"{k}: {v}" for k, v in value.items())
            bits.append(f"{key}: {value}")
        paragraphs.append(". ".join(bits))
    return "\n\n".join(paragraphs)


def extract(path: str) -> Tuple[str, str]:
    """(title, text) for a file. Text is '' when nothing readable was found."""
    lower = path.lower()
    title = os.path.splitext(os.path.basename(path))[0]
    if lower.endswith(PDF_EXTENSIONS):
        text = _read_pdf(path)
    elif lower.endswith(DOCX_EXTENSIONS):
        text = _read_docx(path)
    elif lower.endswith(JSON_EXTENSIONS):
        text = _render_json(path, _read_text(path))
    else:
        text = _read_text(path)
        first = next((line.strip("# ").strip() for line in text.splitlines() if line.strip()), "")
        if lower.endswith((".md", ".markdown")) and first and len(first) < 120:
            title = first
    return title, text


def chunk_text(text: str, size: int = CHUNK_CHARS) -> List[str]:
    """Split on paragraph boundaries, packing paragraphs up to ~size chars."""
    chunks, current = [], ""
    for para in text.split("\n\n"):
        para = para.strip()
        if not para:
            continue
        if len(current) + len(para) + 2 <= size:
            current = f"{current}\n\n{para}" if current else para
            continue
        if current:
            chunks.append(current)
        while len(para) > size:
            chunks.append(para[:size])
            para = para[size:]
        current = para
    if current:
        chunks.append(current)
    return chunks


# ───────────────────────── indexing ─────────────────────────


def _fingerprint(path: str) -> Tuple[float, int]:
    stat = os.stat(path)
    return round(stat.st_mtime, 3), stat.st_size


def _sha1(path: str) -> str:
    digest = hashlib.sha1()
    with open(path, "rb") as handle:
        for block in iter(lambda: handle.read(1 << 16), b""):
            digest.update(block)
    return digest.hexdigest()


def ingest_file(path: str, tag: str, force: bool = False, memory_store=None) -> str:
    """
    Index one file. Returns 'skipped' | 'updated' | 'empty' | 'failed'.
    """
    store = memory_store or _memory()
    try:
        mtime, size = _fingerprint(path)
    except OSError:
        return "failed"
    known = store.ingest_source(path)
    if known and not force and known.get("mtime") == mtime and known.get("size") == size:
        return "skipped"
    try:
        title, text = extract(path)
    except Exception as exc:
        store.upsert_ingest_source(path, tag, mtime, size, "", 0, "failed", str(exc)[:200])
        return "failed"
    minimum = MIN_JSON_CHARS if path.lower().endswith(JSON_EXTENSIONS) else MIN_TEXT_CHARS
    if len(text.strip()) < minimum:
        store.remove_document(path)
        store.upsert_ingest_source(path, tag, mtime, size, "", 0, "empty", "")
        return "empty"
    chunks = chunk_text(text)
    stored = store.add_document(path, chunks, tag=tag, title=title)
    store.upsert_ingest_source(path, tag, mtime, size, _sha1(path), stored, "ok", "")
    return "updated"


def reindex(
    roots: Optional[Iterable[Tuple[str, str]]] = None,
    force: bool = False,
    budget_seconds: Optional[float] = None,
    memory_store=None,
) -> Dict[str, Any]:
    """
    Walk the roots and bring the index up to date. `budget_seconds` stops
    early (used for the startup pass); whatever is left waits for the nightly
    job. Files that vanished since the last pass are removed.
    """
    store = memory_store or _memory()
    started = time.perf_counter()
    stats = {"scanned": 0, "updated": 0, "skipped": 0, "empty": 0, "failed": 0, "removed": 0, "partial": False}
    seen = set()
    roots = list(roots if roots is not None else parse_roots())
    for path, tag in scan(roots):
        stats["scanned"] += 1
        seen.add(os.path.normcase(path))
        if budget_seconds is not None and time.perf_counter() - started > budget_seconds:
            stats["partial"] = True
            break
        outcome = ingest_file(path, tag, force=force, memory_store=store)
        stats[outcome] += 1
    if not stats["partial"]:
        root_dirs = [os.path.normcase(path) for _, path in roots]
        for known in store.ingest_sources():
            known_path = known["path"]
            if os.path.normcase(known_path) in seen:
                continue
            # Only forget files that lived under a root we just walked.
            if any(os.path.normcase(known_path).startswith(r) for r in root_dirs) and not os.path.exists(known_path):
                store.remove_document(known_path)
                store.remove_ingest_source(known_path)
                stats["removed"] += 1
    stats["seconds"] = round(time.perf_counter() - started, 1)
    stats["roots"] = [f"{tag}={path}" for tag, path in roots]
    return stats


def _memory():
    from memory import memory

    return memory
